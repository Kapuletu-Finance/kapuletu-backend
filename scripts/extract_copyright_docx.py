import os
import re
import argparse
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Configuration
COPYRIGHT_YEAR = datetime.now().year
COMPANY_NAME = "Kapuletu Systems"
LOGO_PATH = os.path.join("assets", "logo.jpg")

SECRET_PATTERNS = [
    re.compile(r'(?i)api[_-]?key\s*[:=]\s*[\'"][^\'"]+[\'"]'),
    re.compile(r'(?i)password\s*[:=]\s*[\'"][^\'"]+[\'"]'),
    re.compile(r'(?i)secret\s*[:=]\s*[\'"][^\'"]+[\'"]'),
    re.compile(r'(?i)token\s*[:=]\s*[\'"][^\'"]+[\'"]')
]

def get_file_content_and_check(filepath):
    """Reads a file and flags potential secrets."""
    if not os.path.exists(filepath):
        print(f"[-] Error: File not found -> {filepath}")
        return None

    try:
        with open(filepath, 'r', encoding='utf-8') as f:
            content = f.read()
    except Exception as e:
        print(f"[-] Error reading {filepath}: {e}")
        return None

    flags = []
    for i, line in enumerate(content.split('\n')):
        for pattern in SECRET_PATTERNS:
            if pattern.search(line):
                flags.append((i + 1, line.strip()))

    if flags:
        print(f"[!] WARNING: Potential secrets found in {filepath}:")
        for line_num, line_content in flags:
            print(f"    Line {line_num}: {line_content}")
        print("    Please redact these before submitting!")
    
    return content

def main():
    parser = argparse.ArgumentParser(description="Extract and generate a DOCX for copyright submission.")
    parser.add_argument('-i', '--input', required=True, help="Text file containing a list of file paths to extract (one per line).")
    parser.add_argument('-o', '--output', default='copyright_submission.docx', help="Output DOCX file name (default: copyright_submission.docx)")
    
    args = parser.parse_args()

    if not os.path.exists(args.input):
        print(f"[-] Input list file not found: {args.input}")
        return

    print(f"[*] Reading file list from {args.input}...")
    with open(args.input, 'r', encoding='utf-8') as f:
        file_list = [line.strip() for line in f if line.strip() and not line.startswith('#')]

    if not file_list:
        print("[-] No files to process. Please add file paths to your input list.")
        return
        
    print(f"[*] Found {len(file_list)} files to process.")
    
    doc = Document()
    
    # Title Page
    if os.path.exists(LOGO_PATH):
        try:
            doc.add_picture(LOGO_PATH, width=Inches(2.5))
        except Exception:
            pass

    title = doc.add_heading('Software Source Code Registration', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.runs[0]
    title_run.font.color.rgb = RGBColor(0x1e, 0x3a, 0x8a) # #1e3a8a

    doc.add_paragraph(f"Company: {COMPANY_NAME}")
    doc.add_paragraph(f"Date Generated: {datetime.now().strftime('%B %d, %Y')}")
    
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    runner = p.add_run("CONFIDENTIAL AND PROPRIETARY")
    runner.bold = True
    doc.add_paragraph("This document contains unpublished, proprietary source code of Kapuletu Systems. It is provided strictly for the purpose of copyright registration and may not be reproduced or distributed.")
    
    doc.add_page_break()

    # Code Blocks
    success_count = 0
    for filepath in file_list:
        content = get_file_content_and_check(filepath)
        if content:
            # File Header
            h2 = doc.add_heading(f"File: {filepath}", level=2)
            h2_run = h2.runs[0]
            h2_run.font.color.rgb = RGBColor(0xb9, 0x1c, 0x1c) # #b91c1c
            
            # Code Content
            p = doc.add_paragraph(content)
            # Use Courier for code style
            p.style.font.name = 'Courier'
            p.style.font.size = Pt(8)
            
            doc.add_page_break()
            success_count += 1
            print(f"[+] Processed: {filepath}")

    print(f"[*] Building DOCX (this may take a moment)...")
    try:
        doc.save(args.output)
        print(f"\n[*] Done! Successfully processed {success_count}/{len(file_list)} files.")
        print(f"[*] Output saved to: {args.output}")
    except Exception as e:
        print(f"[-] Failed to generate DOCX: {e}")

if __name__ == "__main__":
    main()
