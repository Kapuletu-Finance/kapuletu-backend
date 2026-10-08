"""
Contribution volume and ledger integrity, platform-wide and read-only.

This is the treasurers' money (group contributions), not Kapuletu revenue. Finance sees how much flows through
the platform and whether the sealed records are intact; it never edits them.
"""
import csv
import datetime
import io
from collections import defaultdict
from decimal import Decimal
from typing import Optional

from sqlalchemy import func, or_
from sqlalchemy.orm import Session

from models.campaign import Campaign
from models.group import Group
from models.transaction import Transaction
from models.users import User
from services.finance.ledger_service import LedgerService

from .common import FinanceError, audit, iso, page_params, paged, user_name

ZERO = Decimal("0")
INTEGRITY_LIMIT = 20000  # records checked per run; narrow the dates for more
EXPORT_LIMIT = 20000
PDF_EXPORT_LIMIT = 2000


def _spreadsheet_safe(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return f"'{value}"
    return value


class VolumeService:
    def __init__(self, db: Session):
        self.db = db

    def _approved(self, start: datetime.datetime, end: datetime.datetime):
        return self.db.query(Transaction).filter(
            Transaction.status == "approved", Transaction.created_at >= start, Transaction.created_at < end
        )

    def summary(self, start: datetime.datetime, end: datetime.datetime) -> dict:
        totals = self.db.query(func.count(Transaction.transaction_id), func.coalesce(func.sum(Transaction.amount), 0),
                               func.count(func.distinct(Transaction.group_id)),
                               func.count(func.distinct(Transaction.owner_id))).filter(
            Transaction.status == "approved", Transaction.created_at >= start, Transaction.created_at < end
        ).one()
        count, amount, groups, treasurers = totals

        by_month = defaultdict(lambda: [0, ZERO])
        by_method = defaultdict(lambda: [0, ZERO])
        for created_at, method, value in self.db.query(Transaction.created_at, Transaction.payment_method,
                                                       Transaction.amount).filter(
            Transaction.status == "approved", Transaction.created_at >= start, Transaction.created_at < end
        ):
            m = by_month[created_at.strftime("%Y-%m")]
            m[0] += 1
            m[1] += Decimal(str(value or 0))
            k = by_method[method or "Unknown"]
            k[0] += 1
            k[1] += Decimal(str(value or 0))

        top = self.db.query(Group.group_id, Group.group_name, Group.owner_id,
                            func.count(Transaction.transaction_id), func.sum(Transaction.amount)).join(
            Transaction, Transaction.group_id == Group.group_id
        ).filter(
            Transaction.status == "approved", Transaction.created_at >= start, Transaction.created_at < end
        ).group_by(Group.group_id, Group.group_name, Group.owner_id).order_by(func.sum(Transaction.amount).desc()).limit(10).all()
        owners = {u.user_id: u for u in self.db.query(User).filter(User.user_id.in_([t[2] for t in top]))} if top else {}

        return {
            "period": {"from": start.isoformat(), "to": end.isoformat()},
            "totals": {
                "contributions": count, "amount": float(amount or 0), "active_groups": groups,
                "active_treasurers": treasurers,
                "average_contribution": round(float(amount or 0) / count, 2) if count else 0.0,
            },
            "by_month": [{"month": k, "contributions": v[0], "amount": float(v[1])} for k, v in sorted(by_month.items())],
            "by_method": [{"method": k, "contributions": v[0], "amount": float(v[1])}
                          for k, v in sorted(by_method.items(), key=lambda x: -x[1][1])],
            "top_groups": [{
                "group_id": str(gid), "group_name": name, "treasurer": user_name(owners.get(owner)),
                "treasurer_id": str(owner), "contributions": n, "amount": float(total or 0),
            } for gid, name, owner, n, total in top],
        }

    def _contribution_query(
        self,
        start: datetime.datetime,
        end: datetime.datetime,
        search: Optional[str] = None,
        group: Optional[str] = None,
        treasurer: Optional[str] = None,
        contributor: Optional[str] = None,
        method: Optional[str] = None,
    ):
        query = self.db.query(Transaction, Group, User, Campaign).join(
            Group, Transaction.group_id == Group.group_id
        ).join(User, Transaction.owner_id == User.user_id).outerjoin(
            Campaign, Transaction.campaign_id == Campaign.campaign_id
        ).filter(
            Transaction.status == "approved",
            Transaction.created_at >= start,
            Transaction.created_at < end,
        )
        if search:
            like = f"%{search.strip()}%"
            query = query.filter(or_(
                Transaction.sender_name.ilike(like),
                Transaction.sender_phone.ilike(like),
                Transaction.transaction_code.ilike(like),
                Group.group_name.ilike(like),
                User.first_name.ilike(like),
                User.last_name.ilike(like),
                User.email.ilike(like),
                Campaign.title.ilike(like),
            ))
        if group:
            query = query.filter(Group.group_name.ilike(f"%{group.strip()}%"))
        if treasurer:
            like = f"%{treasurer.strip()}%"
            query = query.filter(or_(User.first_name.ilike(like), User.last_name.ilike(like),
                                     User.email.ilike(like)))
        if contributor:
            like = f"%{contributor.strip()}%"
            query = query.filter(or_(Transaction.sender_name.ilike(like), Transaction.sender_phone.ilike(like)))
        if method:
            query = query.filter(Transaction.payment_method.ilike(f"%{method.strip()}%"))
        return query

    @staticmethod
    def _contribution_row(transaction: Transaction, group: Group, treasurer: User,
                          campaign: Optional[Campaign]) -> dict:
        return {
            "transaction_id": str(transaction.transaction_id),
            "transaction_code": transaction.transaction_code,
            "created_at": iso(transaction.created_at),
            "contributor_name": transaction.sender_name,
            "contributor_phone": transaction.sender_phone,
            "group_id": str(group.group_id),
            "group_name": group.group_name,
            "treasurer_id": str(treasurer.user_id),
            "treasurer_name": user_name(treasurer),
            "campaign_id": str(campaign.campaign_id) if campaign else None,
            "campaign_name": campaign.title if campaign else None,
            "payment_method": transaction.payment_method or "Unknown",
            "amount": float(transaction.amount or 0),
        }

    def list_contributions(
        self,
        start: datetime.datetime,
        end: datetime.datetime,
        search: Optional[str] = None,
        group: Optional[str] = None,
        treasurer: Optional[str] = None,
        contributor: Optional[str] = None,
        method: Optional[str] = None,
        page: int = 1,
        limit: int = 50,
    ) -> dict:
        query = self._contribution_query(start, end, search, group, treasurer, contributor, method)
        total, amount = query.with_entities(
            func.count(Transaction.transaction_id), func.coalesce(func.sum(Transaction.amount), 0)
        ).one()
        offset, limit = page_params(page, limit)
        rows = query.order_by(Transaction.created_at.desc(), Transaction.transaction_id).offset(offset).limit(limit).all()
        return {
            **paged([self._contribution_row(*row) for row in rows], total, page, limit),
            "total_amount": float(amount or 0),
            "period": {"from": start.isoformat(), "to": end.isoformat()},
        }

    def export_statement(
        self,
        start: datetime.datetime,
        end: datetime.datetime,
        fmt: str,
        search: Optional[str] = None,
        group: Optional[str] = None,
        treasurer: Optional[str] = None,
        contributor: Optional[str] = None,
        method: Optional[str] = None,
        prepared_by: Optional[str] = None,
        actor_id=None,
    ) -> tuple[bytes, str, str]:
        if fmt not in ("csv", "excel", "pdf"):
            raise FinanceError("format must be csv, excel or pdf")
        query = self._contribution_query(start, end, search, group, treasurer, contributor, method)
        limit = PDF_EXPORT_LIMIT if fmt == "pdf" else EXPORT_LIMIT
        rows = query.order_by(Transaction.created_at.asc(), Transaction.transaction_id).limit(limit + 1).all()
        if len(rows) > limit:
            raise FinanceError(
                f"This statement exceeds the {limit:,}-record {fmt.upper()} export limit; narrow the filters"
            )

        records = [self._contribution_row(*row) for row in rows]
        total = sum((Decimal(str(row["amount"])) for row in records), ZERO)
        columns = ["Date", "Contributor", "Phone", "Group", "Treasurer", "Campaign", "Method", "Reference", "Amount (KES)"]
        values = [[
            row["created_at"] or "",
            row["contributor_name"] or "—",
            row["contributor_phone"] or "—",
            row["group_name"],
            row["treasurer_name"],
            row["campaign_name"] or "—",
            row["payment_method"],
            row["transaction_code"],
            row["amount"],
        ] for row in records]
        period_label = f"{start:%d %b %Y} – {(end - datetime.timedelta(seconds=1)):%d %b %Y}"
        stem = f"kapuletu_contributions_{start:%Y%m%d}_{(end - datetime.timedelta(seconds=1)):%Y%m%d}"

        if fmt == "csv":
            output = io.StringIO()
            writer = csv.writer(output)
            writer.writerow(["Kapuletu Contribution Cashflow Statement", period_label])
            writer.writerow(["Approved contributions", len(records), "Total collected (KES)", f"{total:.2f}"])
            writer.writerow([])
            writer.writerow(columns)
            writer.writerows([[_spreadsheet_safe(value) for value in row] for row in values])
            return output.getvalue().encode("utf-8-sig"), "text/csv", f"{stem}.csv"

        if fmt == "excel":
            from openpyxl import Workbook
            from openpyxl.styles import Font

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Contributions"
            sheet.append(["Kapuletu Contribution Cashflow Statement", period_label])
            sheet["A1"].font = Font(bold=True, size=14)
            sheet.append(["Approved contributions", len(records), "Total collected (KES)", float(total)])
            sheet.append([])
            sheet.append(columns)
            for cell in sheet[4]:
                cell.font = Font(bold=True)
            for row in values:
                sheet.append([_spreadsheet_safe(value) for value in row])
            for index, column_cells in enumerate(sheet.columns, start=1):
                width = max(len(str(cell.value or "")) for cell in column_cells)
                sheet.column_dimensions[sheet.cell(1, index).column_letter].width = min(max(12, width + 2), 40)
            sheet.freeze_panes = "A5"
            sheet.auto_filter.ref = f"A4:I{sheet.max_row}"
            output = io.BytesIO()
            workbook.save(output)
            return output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", f"{stem}.xlsx"

        from services.documents.official import OfficialDocument

        document = OfficialDocument(
            self.db,
            title="Contribution Cashflow Statement",
            subtitle=period_label,
            department="FIN",
            doc_type="CFS",
            prepared_by=prepared_by,
            orientation="landscape",
        )
        document.paragraph(
            f"{len(records):,} approved contributions · KES {total:,.2f} collected. "
            "These funds belong to treasurers' groups and are not Kapuletu revenue.",
            muted=True,
        )
        document.table(
            columns,
            [[*row[:-1], f"{row[-1]:,.2f}"] for row in values],
            col_widths=[0.11, 0.12, 0.10, 0.13, 0.12, 0.11, 0.09, 0.13, 0.09],
            numeric_cols={8},
        )
        document.signature_block(note="Statement includes approved contribution records matching the selected filters.")
        content = document.build(
            actor_id=str(actor_id) if actor_id else None,
            audit_details={"report": "contribution_cashflow", "period": period_label,
                           "records": len(records), "total": str(total)},
        )
        return content, "application/pdf", f"{document.filename_stem}.pdf"

    def integrity(self, start: datetime.datetime, end: datetime.datetime, actor_id=None) -> dict:
        """Re-computes the SHA-256 seal of every approved contribution in the period."""
        hasher = LedgerService(self.db)
        checked = unsealed = 0
        tampered = []
        txns = self._approved(start, end).order_by(Transaction.created_at).limit(INTEGRITY_LIMIT + 1).all()
        truncated = len(txns) > INTEGRITY_LIMIT
        for txn in txns[:INTEGRITY_LIMIT]:
            checked += 1
            if not txn.ledger_hash:
                unsealed += 1
                continue
            if hasher._recalculate_hash(txn) != txn.ledger_hash:
                tampered.append(txn)

        groups = {g.group_id: g.group_name for g in self.db.query(Group).filter(
            Group.group_id.in_([t.group_id for t in tampered]))} if tampered else {}
        audit(self.db, actor_id, "LEDGER_INTEGRITY_CHECK", "LEDGER", None, {
            "from": iso(start), "to": iso(end), "checked": checked, "tampered": len(tampered), "unsealed": unsealed,
        })
        self.db.commit()
        return {
            "period": {"from": start.isoformat(), "to": end.isoformat()},
            "checked": checked,
            "intact": checked - unsealed - len(tampered),
            "unsealed": unsealed,
            "truncated": truncated,
            "tampered": [{
                "transaction_id": str(t.transaction_id), "group_name": groups.get(t.group_id),
                "transaction_code": t.transaction_code, "amount": float(t.amount or 0), "created_at": iso(t.created_at),
            } for t in tampered],
        }
