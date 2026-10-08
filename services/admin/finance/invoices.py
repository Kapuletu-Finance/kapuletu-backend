"""Invoices and payments: filtered lists and detail views."""
import csv
import datetime
import io
from decimal import Decimal
from typing import Optional

from sqlalchemy import or_
from sqlalchemy.orm import Session

from models.billing import CreditNote, Invoice, InvoiceLine, ProviderEvent, Refund
from models.subscription import Plan, Subscription, SubscriptionPayment
from models.users import User

from .common import FinanceError, as_uuid, iso, page_params, paged, user_name

INVOICE_STATUSES = ("draft", "open", "paid", "void")
PAYMENT_STATUSES = ("pending", "initiated", "success", "failed")
PAYMENT_TYPES = ("payment", "refund", "comp")
EXPORT_LIMIT = 20000
PDF_EXPORT_LIMIT = 2000
EXCEL_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _spreadsheet_safe(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return f"'{value}"
    return value


def _user_search(q: str):
    like = f"%{q.strip()}%"
    return or_(User.first_name.ilike(like), User.last_name.ilike(like), User.email.ilike(like),
               User.phone_number.ilike(like), User.slug.ilike(like))


def money_out(value) -> float:
    return float(value or 0)


class InvoiceService:
    def __init__(self, db: Session):
        self.db = db

    # --- invoices ---

    @staticmethod
    def _invoice_row(inv: Invoice, user: Optional[User]) -> dict:
        return {
            "invoice_id": str(inv.invoice_id),
            "number": inv.number,
            "user_id": str(inv.user_id),
            "user_name": user_name(user),
            "status": inv.status,
            "currency": inv.currency,
            "subtotal": money_out(inv.subtotal),
            "tax": money_out(inv.tax),
            "total": money_out(inv.total),
            "billing_cycle": inv.billing_cycle,
            "period_start": iso(inv.period_start),
            "period_end": iso(inv.period_end),
            "issued_at": iso(inv.issued_at),
            "paid_at": iso(inv.paid_at),
        }

    def _invoice_query(self, status, q, date_from, date_to, user_id):
        query = self.db.query(Invoice, User).join(User, Invoice.user_id == User.user_id)
        if status:
            if status not in INVOICE_STATUSES:
                raise FinanceError(f"status must be one of {', '.join(INVOICE_STATUSES)}")
            query = query.filter(Invoice.status == status)
        if user_id:
            query = query.filter(Invoice.user_id == as_uuid(user_id))
        if date_from:
            query = query.filter(Invoice.issued_at >= date_from)
        if date_to:
            query = query.filter(Invoice.issued_at < date_to)
        if q:
            query = query.filter(or_(_user_search(q), Invoice.number.ilike(f"%{q.strip()}%")))
        return query

    def list_invoices(self, status: Optional[str] = None, q: Optional[str] = None,
                      date_from: Optional[datetime.datetime] = None, date_to: Optional[datetime.datetime] = None,
                      user_id: Optional[str] = None, page: int = 1, limit: int = 50) -> dict:
        query = self._invoice_query(status, q, date_from, date_to, user_id)
        total = query.count()
        offset, limit = page_params(page, limit)
        rows = query.order_by(Invoice.issued_at.desc()).offset(offset).limit(limit).all()
        return paged([self._invoice_row(i, u) for i, u in rows], total, page, limit)

    def _payment_query(self, status, provider, type_, q, date_from, date_to, user_id):
        query = self.db.query(SubscriptionPayment, User, Plan, Invoice.number).join(
            User, SubscriptionPayment.user_id == User.user_id
        ).outerjoin(Subscription, SubscriptionPayment.subscription_id == Subscription.subscription_id).outerjoin(
            Plan, Subscription.plan_id == Plan.plan_id
        ).outerjoin(Invoice, SubscriptionPayment.invoice_id == Invoice.invoice_id)
        if status:
            if status not in PAYMENT_STATUSES:
                raise FinanceError(f"status must be one of {', '.join(PAYMENT_STATUSES)}")
            query = query.filter(SubscriptionPayment.status == status)
        if provider:
            query = query.filter(SubscriptionPayment.payment_method == provider)
        if type_:
            if type_ not in PAYMENT_TYPES:
                raise FinanceError(f"type must be one of {', '.join(PAYMENT_TYPES)}")
            if type_ == "payment":
                query = query.filter(or_(SubscriptionPayment.transaction_type.is_(None),
                                         SubscriptionPayment.transaction_type == "payment"))
            else:
                query = query.filter(SubscriptionPayment.transaction_type == type_)
        if user_id:
            query = query.filter(SubscriptionPayment.user_id == as_uuid(user_id))
        if date_from:
            query = query.filter(SubscriptionPayment.created_at >= date_from)
        if date_to:
            query = query.filter(SubscriptionPayment.created_at < date_to)
        if q:
            like = f"%{q.strip()}%"
            query = query.filter(or_(_user_search(q), SubscriptionPayment.provider_reference.ilike(like),
                                     Invoice.number.ilike(like)))
        return query

    @staticmethod
    def _export_file(
        *,
        title: str,
        period: str,
        columns: list[str],
        rows: list[list],
        fmt: str,
        filename_stem: str,
        summary: list[tuple[str, str]],
        pdf_columns: list[str],
        pdf_rows: list[list],
        db: Session,
        prepared_by: Optional[str],
        actor_id,
    ) -> tuple[bytes, str, str]:
        if fmt not in ("csv", "excel", "pdf"):
            raise FinanceError("format must be csv, excel or pdf")

        if fmt == "csv":
            output = io.StringIO()
            writer = csv.writer(output)
            writer.writerow([title, period])
            writer.writerows(summary)
            writer.writerow([])
            writer.writerow(columns)
            writer.writerows([[_spreadsheet_safe(value) for value in row] for row in rows])
            return output.getvalue().encode("utf-8-sig"), "text/csv", f"{filename_stem}.csv"

        if fmt == "excel":
            from openpyxl import Workbook
            from openpyxl.styles import Font

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Records"
            sheet.append([title, period])
            sheet["A1"].font = Font(bold=True, size=14)
            for item in summary:
                sheet.append(list(item))
            sheet.append([])
            header_row = sheet.max_row + 1
            sheet.append(columns)
            for cell in sheet[header_row]:
                cell.font = Font(bold=True)
            for row in rows:
                sheet.append([_spreadsheet_safe(value) for value in row])
            sheet.freeze_panes = f"A{header_row + 1}"
            sheet.auto_filter.ref = f"A{header_row}:{sheet.cell(sheet.max_row, len(columns)).coordinate}"
            for column_cells in sheet.columns:
                width = max(len(str(cell.value or "")) for cell in column_cells)
                sheet.column_dimensions[column_cells[0].column_letter].width = min(max(12, width + 2), 40)
            output = io.BytesIO()
            workbook.save(output)
            return output.getvalue(), EXCEL_MIME, f"{filename_stem}.xlsx"

        from services.documents.official import OfficialDocument

        document = OfficialDocument(
            db,
            title=title,
            subtitle=period,
            department="FIN",
            doc_type="PAY" if "Payment" in title else "INV",
            prepared_by=prepared_by,
            orientation="landscape",
        )
        document.key_figures([(label, value, None) for label, value in summary[:4]])
        document.table(pdf_columns, pdf_rows, empty_message="No matching records.")
        document.signature_block(note="Includes all records matching the selected filters.")
        content = document.build(
            actor_id=str(actor_id) if actor_id else None,
            audit_details={"report": filename_stem, "records": len(rows), "period": period},
        )
        return content, "application/pdf", f"{document.filename_stem}.pdf"

    @staticmethod
    def _export_period(date_from, date_to) -> tuple[str, str]:
        if date_from and date_to:
            period = f"{date_from:%d %b %Y} – {(date_to - datetime.timedelta(seconds=1)):%d %b %Y}"
            suffix = f"{date_from:%Y%m%d}_{(date_to - datetime.timedelta(seconds=1)):%Y%m%d}"
        elif date_from:
            period, suffix = f"From {date_from:%d %b %Y}", f"from_{date_from:%Y%m%d}"
        elif date_to:
            period, suffix = f"Through {(date_to - datetime.timedelta(seconds=1)):%d %b %Y}", f"to_{(date_to - datetime.timedelta(seconds=1)):%Y%m%d}"
        else:
            period, suffix = "All dates", "all_dates"
        return period, suffix

    def export_invoices(self, fmt: str, status: Optional[str] = None, q: Optional[str] = None,
                        date_from: Optional[datetime.datetime] = None, date_to: Optional[datetime.datetime] = None,
                        user_id: Optional[str] = None, prepared_by: Optional[str] = None, actor_id=None
                        ) -> tuple[bytes, str, str]:
        if fmt not in ("csv", "excel", "pdf"):
            raise FinanceError("format must be csv, excel or pdf")
        query = self._invoice_query(status, q, date_from, date_to, user_id)
        limit = PDF_EXPORT_LIMIT if fmt == "pdf" else EXPORT_LIMIT
        records = query.order_by(Invoice.issued_at.desc()).limit(limit + 1).all()
        if len(records) > limit:
            raise FinanceError(f"This export exceeds the {limit:,}-record {fmt.upper()} limit; narrow the filters")
        rows = [[
            inv.number, inv.issued_at.isoformat() if inv.issued_at else "", user_name(user), inv.status,
            inv.currency, money_out(inv.subtotal), money_out(inv.tax), money_out(inv.total),
            inv.billing_cycle or "", iso(inv.period_start) or "", iso(inv.period_end) or "",
            iso(inv.paid_at) or "", str(inv.user_id), str(inv.invoice_id),
        ] for inv, user in records]
        period, suffix = self._export_period(date_from, date_to)
        total = sum((Decimal(str(row[7] or 0)) for row in rows), Decimal("0"))
        paid_total = sum((Decimal(str(row[7] or 0)) for row in rows if row[3] == "paid"), Decimal("0"))
        summary = [
            ("Matching invoices", str(len(rows))),
            ("Invoice total (KES)", f"{total:.2f}"),
            ("Paid invoice total (KES)", f"{paid_total:.2f}"),
        ]
        pdf_columns = ["Invoice", "Issued", "Customer", "Status", "Currency", "Subtotal", "Tax", "Total", "Cycle"]
        pdf_rows = [[row[i] for i in (0, 1, 2, 3, 4, 5, 6, 7, 8)] for row in rows]
        return self._export_file(
            title="Kapuletu Invoice Register",
            period=period,
            columns=["Invoice", "Issued at", "Customer", "Status", "Currency", "Subtotal", "Tax", "Total",
                     "Billing cycle", "Period start", "Period end", "Paid at", "Customer ID", "Invoice ID"],
            rows=rows,
            fmt=fmt,
            filename_stem=f"kapuletu_invoices_{suffix}",
            summary=summary,
            pdf_columns=pdf_columns,
            pdf_rows=pdf_rows,
            db=self.db,
            prepared_by=prepared_by,
            actor_id=actor_id,
        )

    def get_invoice(self, invoice_id: str) -> dict:
        inv = self.db.get(Invoice, as_uuid(invoice_id))
        if not inv:
            raise FinanceError("Invoice not found", 404)
        out = self._invoice_row(inv, self.db.get(User, inv.user_id))
        out["notes"] = inv.notes
        out["lines"] = [{
            "kind": l.kind, "description": l.description, "quantity": l.quantity,
            "unit_amount": money_out(l.unit_amount), "amount": money_out(l.amount),
        } for l in self.db.query(InvoiceLine).filter(InvoiceLine.invoice_id == inv.invoice_id).all()]
        out["payments"] = [self._payment_row(p, None, None, inv.number) for p in self.db.query(SubscriptionPayment).filter(
            SubscriptionPayment.invoice_id == inv.invoice_id).order_by(SubscriptionPayment.created_at).all()]
        out["credit_notes"] = [{
            "number": c.number, "amount": money_out(c.amount), "reason": c.reason, "issued_at": iso(c.issued_at),
        } for c in self.db.query(CreditNote).filter(CreditNote.invoice_id == inv.invoice_id).all()]
        return out

    def invoice_pdf(self, invoice_id: str) -> tuple[bytes, str]:
        inv = self.db.get(Invoice, as_uuid(invoice_id))
        if not inv:
            raise FinanceError("Invoice not found", 404)
        lines = self.db.query(InvoiceLine).filter(InvoiceLine.invoice_id == inv.invoice_id).all()
        credit_notes = self.db.query(CreditNote).filter(CreditNote.invoice_id == inv.invoice_id).all()
        user = self.db.get(User, inv.user_id)

        from services.finance.receipt_document import render_invoice

        return render_invoice(self.db, inv, lines, user, credit_notes)

    # --- payments ---

    @staticmethod
    def _payment_row(p: SubscriptionPayment, user: Optional[User], plan: Optional[Plan], invoice_number) -> dict:
        meta = p.payment_metadata or {}
        return {
            "payment_id": str(p.payment_id),
            "user_id": str(p.user_id),
            "user_name": user_name(user) if user is not None else None,
            "plan_name": plan.name if plan else None,
            "amount": money_out(p.amount),
            "currency": p.currency,
            "status": p.status,
            "method": p.payment_method,
            "transaction_type": p.transaction_type or "payment",
            "invoice_number": invoice_number,
            "receipt_number": meta.get("receipt_number"),
            "provider_reference": p.provider_reference,
            "failure_reason": meta.get("failure_reason"),
            "refunded": bool(meta.get("refund_payment_id")),
            "refund_id": meta.get("refund_id"),
            "created_at": iso(p.created_at),
        }

    def list_payments(self, status: Optional[str] = None, provider: Optional[str] = None, type_: Optional[str] = None,
                      q: Optional[str] = None, date_from: Optional[datetime.datetime] = None,
                      date_to: Optional[datetime.datetime] = None, user_id: Optional[str] = None,
                      page: int = 1, limit: int = 50) -> dict:
        query = self._payment_query(status, provider, type_, q, date_from, date_to, user_id)
        total = query.count()
        offset, limit = page_params(page, limit)
        rows = query.order_by(SubscriptionPayment.created_at.desc()).offset(offset).limit(limit).all()
        return paged([self._payment_row(p, u, pl, n) for p, u, pl, n in rows], total, page, limit)

    def export_payments(self, fmt: str, status: Optional[str] = None, provider: Optional[str] = None,
                        type_: Optional[str] = None, q: Optional[str] = None,
                        date_from: Optional[datetime.datetime] = None, date_to: Optional[datetime.datetime] = None,
                        user_id: Optional[str] = None, prepared_by: Optional[str] = None, actor_id=None
                        ) -> tuple[bytes, str, str]:
        if fmt not in ("csv", "excel", "pdf"):
            raise FinanceError("format must be csv, excel or pdf")
        query = self._payment_query(status, provider, type_, q, date_from, date_to, user_id)
        limit = PDF_EXPORT_LIMIT if fmt == "pdf" else EXPORT_LIMIT
        records = query.order_by(SubscriptionPayment.created_at.desc()).limit(limit + 1).all()
        if len(records) > limit:
            raise FinanceError(f"This export exceeds the {limit:,}-record {fmt.upper()} limit; narrow the filters")
        payment_rows = [self._payment_row(p, user, plan, invoice) for p, user, plan, invoice in records]
        rows = [[
            row["created_at"] or "", row["user_name"] or "—", row["plan_name"] or "—", row["transaction_type"],
            row["status"], row["method"] or "—", row["invoice_number"] or "—", row["receipt_number"] or "—",
            row["provider_reference"] or "—", row["amount"], row["currency"], row["failure_reason"] or "—",
            "Yes" if row["refunded"] else "No", row["refund_id"] or "—", row["user_id"], row["payment_id"],
        ] for row in payment_rows]
        period, suffix = self._export_period(date_from, date_to)
        paid_total = sum((Decimal(str(row["amount"] or 0)) for row in payment_rows
                          if row["status"] == "success" and row["transaction_type"] == "payment"), Decimal("0"))
        refund_total = sum((Decimal(str(row["amount"] or 0)) for row in payment_rows
                            if row["transaction_type"] == "refund"), Decimal("0"))
        summary = [
            ("Matching payment records", str(len(rows))),
            ("Successful payments (KES)", f"{paid_total:.2f}"),
            ("Refund entries (KES)", f"{refund_total:.2f}"),
        ]
        pdf_columns = ["Date", "Customer", "Plan", "Type", "Status", "Method", "Invoice", "Receipt", "Amount", "Currency"]
        pdf_rows = [[row[i] for i in (0, 1, 2, 3, 4, 5, 6, 7, 9, 10)] for row in rows]
        return self._export_file(
            title="Kapuletu Payments Register",
            period=period,
            columns=["Date", "Customer", "Plan", "Type", "Status", "Method", "Invoice", "Receipt",
                     "Provider reference", "Amount", "Currency", "Failure reason", "Refunded", "Refund ID",
                     "Customer ID", "Payment ID"],
            rows=rows,
            fmt=fmt,
            filename_stem=f"kapuletu_payments_{suffix}",
            summary=summary,
            pdf_columns=pdf_columns,
            pdf_rows=pdf_rows,
            db=self.db,
            prepared_by=prepared_by,
            actor_id=actor_id,
        )

    def get_payment(self, payment_id: str) -> dict:
        p = self.db.get(SubscriptionPayment, as_uuid(payment_id))
        if not p:
            raise FinanceError("Payment not found", 404)
        sub = self.db.get(Subscription, p.subscription_id) if p.subscription_id else None
        invoice = self.db.get(Invoice, p.invoice_id) if p.invoice_id else None
        out = self._payment_row(p, self.db.get(User, p.user_id), self.db.get(Plan, sub.plan_id) if sub else None,
                                invoice.number if invoice else None)
        out["invoice_id"] = str(invoice.invoice_id) if invoice else None
        out["metadata"] = {k: v for k, v in (p.payment_metadata or {}).items() if k not in ("phone_number", "email")}
        refund = self.db.query(Refund).filter(Refund.payment_id == p.payment_id).order_by(Refund.requested_at.desc()).first()
        out["refund"] = {"refund_id": str(refund.refund_id), "status": refund.status, "amount": money_out(refund.amount)} if refund else None
        out["provider_events"] = [{
            "provider": e.provider, "outcome": e.outcome, "received_at": iso(e.received_at),
            "processed_at": iso(e.processed_at), "payload": e.payload,
        } for e in self.db.query(ProviderEvent).filter(
            ProviderEvent.correlation_id == p.provider_reference).order_by(ProviderEvent.received_at).all()
        ] if p.provider_reference else []
        return out
