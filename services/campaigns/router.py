from typing import List, Dict, Any
from uuid import UUID
from fastapi import APIRouter, Depends, HTTPException, status, Query, Request
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import select, func, cast, String
import random
import io

from common.database import get_db
from common.config import get_config
from common.auth_dependencies import get_verified_user
from services.campaigns.schemas import CampaignCreate, CampaignUpdate, CampaignOut, PaginatedCampaignResponse, PaginatedTransactionResponse, CampaignActivity, ChartDataPoint, CampaignReportPreview, PinResponse
from services.auth.router import limiter
from repositories import campaign_repo, group_repo
from models.audit_log import AuditLog
from models.transaction import Transaction
from services.audit.service import AuditService

router = APIRouter(prefix="", tags=["5. Campaigns Management"])

def _verify_group_ownership(db: Session, group_id: str, owner_id: str):
    """Helper to verify that the group exists and belongs to the current user."""
    group = group_repo.get_group(db=db, identifier=str(group_id))
    if not group:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Group not found.")
    if str(group.owner_id) != owner_id:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="You do not have permission to manage campaigns for this group.")
    return group


@router.post("/groups/{group_id}/campaigns", response_model=CampaignOut, status_code=status.HTTP_201_CREATED, summary="Create Campaign")
@limiter.limit("50/minute")
async def create_campaign(
    request: Request,
    group_id: UUID,
    payload: CampaignCreate,
    db: Session = Depends(get_db),
    current_user: Dict[str, Any] = Depends(get_verified_user)
):
    """Creates a new campaign for a specific group."""
    _verify_group_ownership(db, str(group_id), current_user.get('sub'))
    
    new_campaign = campaign_repo.create_campaign(
        db=db,
        group_id=str(group_id),
        title=payload.title,
        description=payload.description,
        target_amount=payload.target_amount,
        payment_instructions=payload.payment_instructions
    )
    
    AuditService(db).log_action(
        actor_id=current_user.get('sub'),
        action="CAMPAIGN_CREATED",
        entity_type="campaign",
        entity_id=str(new_campaign.campaign_id),
        details={"message": f"New campaign \"{new_campaign.title}\" created", "campaign_id": str(new_campaign.campaign_id)}
    )
    
    return new_campaign

@router.get("/groups/{group_id}/campaigns", response_model=PaginatedCampaignResponse, summary="List Campaigns")
@limiter.limit("50/minute")
async def list_campaigns(
    request: Request,
    group_id: str,
    skip: int = Query(0, ge=0),
    limit: int = Query(10, ge=1, le=100),
    search: str = Query(None, description="Search by title"),
    campaign_status: str = Query(None, description="active, archived, or all"),
    db: Session = Depends(get_db),
    current_user: Dict[str, Any] = Depends(get_verified_user)
):
    """Lists all campaigns for a specific group with pagination, search, and dynamic stats."""
    _verify_group_ownership(db, str(group_id), current_user.get('sub'))
    
    campaigns = campaign_repo.get_group_campaigns(db=db, group_id=str(group_id), skip=skip, limit=limit, search=search, status=campaign_status)
    return campaigns

@router.get("/campaigns/{campaign_id}", response_model=CampaignOut, summary="Get Campaign")
async def get_campaign(
    campaign_id: str,
    db: Session = Depends(get_db),
    current_user: Dict[str, Any] = Depends(get_verified_user)
):
    """Retrieves details of a specific campaign."""
    campaign = campaign_repo.get_campaign(db=db, identifier=str(campaign_id))
    if not campaign:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Campaign not found.")
        
    _verify_group_ownership(db, str(campaign.group_id), current_user.get('sub'))
    return campaign

@router.patch("/campaigns/{campaign_id}", response_model=CampaignOut, summary="Update Campaign")
async def update_campaign(
    campaign_id: str,
    payload: CampaignUpdate,
    db: Session = Depends(get_db),
    current_user: Dict[str, Any] = Depends(get_verified_user)
):
    """Updates campaign details."""
    campaign = campaign_repo.get_campaign(db=db, identifier=str(campaign_id))
    if not campaign:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Campaign not found.")
        
    _verify_group_ownership(db, str(campaign.group_id), current_user.get('sub'))
    
    if not campaign.is_active:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Cannot modify an archived campaign.")
        
    updates = {k: v for k, v in payload.model_dump(exclude_unset=True).items() if v is not None}
    
    if "settings" in updates:
        settings_update = updates.pop("settings")
        current_settings = campaign.settings_override or {}
        current_settings.update(settings_update)
        updates["settings_override"] = current_settings
        
    if not updates:
        return campaign
        
    updated_campaign = campaign_repo.update_campaign(db=db, campaign_id=str(campaign.campaign_id), updates=updates)
    
    AuditService(db).log_action(
        actor_id=current_user.get('sub'),
        action="CAMPAIGN_UPDATED",
        entity_type="campaign",
        entity_id=str(campaign_id),
        details={"message": f"Campaign \"{updated_campaign.title}\" updated", "campaign_id": str(campaign_id)}
    )
    
    return updated_campaign

@router.patch("/campaigns/{campaign_id}/favorite", response_model=CampaignOut, summary="Toggle Favorite Campaign")
async def toggle_favorite_campaign(
    campaign_id: str,
    db: Session = Depends(get_db),
    current_user: Dict[str, Any] = Depends(get_verified_user)
):
    """Toggles the is_favorite status of a campaign."""
    campaign = campaign_repo.get_campaign(db=db, identifier=str(campaign_id))
    if not campaign:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Campaign not found.")
        
    _verify_group_ownership(db, str(campaign.group_id), current_user.get('sub'))
    
    updated_campaign = campaign_repo.update_campaign(db=db, campaign_id=str(campaign.campaign_id), updates={"is_favorite": not campaign.is_favorite})
    return updated_campaign

@router.delete("/campaigns/{campaign_id}", response_model=CampaignOut, summary="Archive Campaign")
async def archive_campaign(
    campaign_id: str,
    db: Session = Depends(get_db),
    current_user: Dict[str, Any] = Depends(get_verified_user)
):
    """Safely archives (soft deletes) a campaign."""
    campaign = campaign_repo.get_campaign(db=db, identifier=str(campaign_id))
    if not campaign:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Campaign not found.")
        
    _verify_group_ownership(db, str(campaign.group_id), current_user.get('sub'))
    
    if not campaign.is_active:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Campaign is already archived.")
        
    archived_campaign = campaign_repo.archive_campaign(db=db, campaign_id=str(campaign.campaign_id))
    
    AuditService(db).log_action(
        actor_id=current_user.get('sub'),
        action="CAMPAIGN_ARCHIVED",
        entity_type="campaign",
        entity_id=str(campaign_id),
        details={"message": f"Campaign \"{archived_campaign.title}\" archived", "campaign_id": str(campaign_id)}
    )
    
    return archived_campaign

@router.post("/campaigns/{campaign_id}/regenerate-pin", response_model=PinResponse, summary="Regenerate Access PIN")
async def regenerate_campaign_pin(
    campaign_id: str,
    db: Session = Depends(get_db),
    current_user: Dict[str, Any] = Depends(get_verified_user)
):
    campaign = campaign_repo.get_campaign(db=db, identifier=str(campaign_id))
    if not campaign:
        raise HTTPException(status_code=404, detail="Campaign not found.")
    _verify_group_ownership(db, str(campaign.group_id), current_user.get('sub'))
    
    new_pin = str(random.randint(1000, 9999))
    current_settings = campaign.settings_override or {}
    current_settings["access_pin"] = new_pin
    
    updated_campaign = campaign_repo.update_campaign(db=db, campaign_id=str(campaign.campaign_id), updates={"settings_override": current_settings})
    return {"pin": new_pin}

@router.get("/campaigns/{campaign_id}/chart-data", response_model=List[ChartDataPoint], summary="Get Contribution Chart Data")
async def get_campaign_chart_data(
    campaign_id: str,
    filter: str = Query("this_month", description="this_week, this_month, this_year, all_time"),
    db: Session = Depends(get_db),
    current_user: Dict[str, Any] = Depends(get_verified_user)
):
    campaign = campaign_repo.get_campaign(db=db, identifier=str(campaign_id))
    if not campaign:
        raise HTTPException(status_code=404, detail="Campaign not found.")
    _verify_group_ownership(db, str(campaign.group_id), current_user.get('sub'))
    
    transactions = db.execute(
        select(Transaction).where(Transaction.campaign_id == str(campaign.campaign_id), Transaction.status == "approved")
        .order_by(Transaction.created_at.asc())
    ).scalars().all()
    
    grouped = {}
    for txn in transactions:
        date_str = txn.created_at.strftime("%Y-%m-%d") if filter in ["this_week", "this_month"] else txn.created_at.strftime("%Y-%m")
        grouped[date_str] = grouped.get(date_str, 0.0) + float(txn.amount)
        
    return [{"date": k, "amount": v} for k, v in grouped.items()]

@router.get("/campaigns/{campaign_id}/transactions", response_model=PaginatedTransactionResponse, summary="List Campaign Transactions")
async def get_campaign_transactions(
    campaign_id: str,
    skip: int = Query(0, ge=0),
    limit: int = Query(10, ge=1, le=100),
    search: str = Query(None),
    db: Session = Depends(get_db),
    current_user: Dict[str, Any] = Depends(get_verified_user)
):
    campaign = campaign_repo.get_campaign(db=db, identifier=str(campaign_id))
    if not campaign:
        raise HTTPException(status_code=404, detail="Campaign not found.")
    _verify_group_ownership(db, str(campaign.group_id), current_user.get('sub'))
    
    query = db.query(Transaction).filter(Transaction.campaign_id == str(campaign.campaign_id), Transaction.status == "approved")
    if search:
        query = query.filter(Transaction.sender_name.ilike(f"%{search}%"))
        
    total_items = query.count()
    transactions = query.order_by(Transaction.created_at.desc()).offset(skip).limit(limit).all()
    
    return {
        "items": transactions,
        "total_items": total_items,
        "total_pages": (total_items + limit - 1) // limit if limit > 0 else 0,
        "page": (skip // limit) + 1 if limit > 0 else 1,
        "limit": limit
    }

@router.get("/campaigns/{campaign_id}/activities", response_model=List[CampaignActivity], summary="Get Campaign Activities")
async def get_campaign_activities(
    campaign_id: str,
    db: Session = Depends(get_db),
    current_user: Dict[str, Any] = Depends(get_verified_user)
):
    campaign = campaign_repo.get_campaign(db=db, identifier=str(campaign_id))
    if not campaign:
        raise HTTPException(status_code=404, detail="Campaign not found.")
    _verify_group_ownership(db, str(campaign.group_id), current_user.get('sub'))
    
    logs = db.execute(
        select(AuditLog).where(
            (
                (AuditLog.entity_type == "campaign") & 
                (AuditLog.entity_id == str(campaign.campaign_id))
            ) | (
                cast(AuditLog.details["campaign_id"], String).ilike(f'%{str(campaign.campaign_id)}%')
            )
        ).order_by(AuditLog.created_at.desc()).limit(10)
    ).scalars().all()
    
    return logs

@router.get("/campaigns/{campaign_id}/report-preview", response_model=CampaignReportPreview, summary="Generate WhatsApp Preview")
async def get_campaign_report_preview(
    campaign_id: str,
    db: Session = Depends(get_db),
    current_user: Dict[str, Any] = Depends(get_verified_user)
):
    campaign = campaign_repo.get_campaign(db=db, identifier=str(campaign_id))
    if not campaign:
        raise HTTPException(status_code=404, detail="Campaign not found.")
    _verify_group_ownership(db, str(campaign.group_id), current_user.get('sub'))
    
    settings = campaign.settings_override or {}
    title = settings.get("report_title", "Campaign Update")
    footer = settings.get("report_footer", "")
    indicator = settings.get("paid_indicator", "✔")
    
    # Calculate raised
    raised = db.query(func.sum(Transaction.amount)).filter(Transaction.campaign_id == str(campaign.campaign_id), Transaction.status == "approved").scalar() or 0.0
    
    transactions = db.query(Transaction).filter(Transaction.campaign_id == str(campaign.campaign_id), Transaction.status == "approved").order_by(Transaction.created_at.desc()).limit(10).all()
    
    lines = []
    lines.append(f"*{title}*")
    if campaign.description:
        lines.append(campaign.description)
    lines.append("")
    lines.append(f"Raised so far: Ksh {raised:,.2f} of Ksh {float(campaign.target_amount):,.2f}")
    if campaign.payment_instructions:
        lines.append(f"{campaign.payment_instructions}")
    lines.append("")
    
    for i, txn in enumerate(transactions, 1):
        name = txn.sender_name or "Anonymous"
        lines.append(f"{i}. {name} - Ksh {float(txn.amount):,.2f} {indicator}")
        
    try:
        blank_slots = int(settings.get("blank_slots", 3))
    except (ValueError, TypeError):
        blank_slots = 3
        
    start_idx = len(transactions) + 1
    for i in range(blank_slots):
        lines.append(f"{start_idx + i}.")
        
    lines.append("")
    remaining = max(0, float(campaign.target_amount) - float(raised))
    if footer:
        lines.append(footer)
    else:
        lines.append(f"We still need Ksh {remaining:,.2f} to reach our goal. Every contribution counts.")
        
    frontend_url = get_config().FRONTEND_URL.rstrip('/')
    public_url = f"{frontend_url}/report/{campaign.slug}"
    
    lines.append(f"View the full report at: {public_url}")
    if not settings.get("remove_watermark", False):
        lines.append("\n*Powered by KapuLetu*")
        
    return {
        "preview_text": "\n".join(lines),
        "title": title,
        "description": campaign.description,
        "raised": float(raised),
        "target": float(campaign.target_amount),
        "contributors": [{"name": txn.sender_name or "Anonymous", "amount": float(txn.amount)} for txn in transactions],
        "payment_instructions": campaign.payment_instructions,
        "footer": footer,
        "public_url": public_url
    }

@router.get("/campaigns/{campaign_id}/export/excel", responses={200: {"content": {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": {}}}}, summary="Export Transactions to Excel")
async def export_campaign_excel(
    campaign_id: str,
    db: Session = Depends(get_db),
    current_user: Dict[str, Any] = Depends(get_verified_user)
):
    campaign = campaign_repo.get_campaign(db=db, identifier=str(campaign_id))
    if not campaign:
        raise HTTPException(status_code=404, detail="Campaign not found.")
    _verify_group_ownership(db, str(campaign.group_id), current_user.get('sub'))
    
    settings = campaign.settings_override or {}
    
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl is not installed. Please install it to export Excel files.")
        
    transactions = db.query(Transaction).filter(Transaction.campaign_id == str(campaign.campaign_id), Transaction.status == "approved").order_by(Transaction.created_at.desc()).all()
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Contributions"
    
    # Header Branding
    start_row = 1
    if not settings.get("remove_watermark", False):
        ws.append(["KapuLetu Campaigns Report"])
        ws.cell(row=1, column=1).font = Font(bold=True, color="1A5D1A", size=14)
        ws.append([])
        start_row = 3
    
    headers = ["Date", "Name", "Phone", "Amount (KES)", "Payment Method"]
    ws.append(headers)
    
    # Style the headers
    green_fill = PatternFill(start_color="1A5D1A", end_color="1A5D1A", fill_type="solid")
    for cell in ws[start_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = green_fill
        
    for txn in transactions:
        ws.append([
            txn.created_at.strftime("%Y-%m-%d %H:%M"),
            txn.sender_name or "Anonymous",
            txn.sender_phone or "",
            float(txn.amount),
            txn.payment_method
        ])
        
    # Footer Branding
    if not settings.get("remove_watermark", False):
        ws.append([])
        ws.append(["Report securely generated by KapuLetu"])
        ws.cell(row=ws.max_row, column=1).font = Font(italic=True, color="808080")
        
    # Auto-adjust column widths
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column].width = adjusted_width
        
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=campaign_{campaign.slug}_contributions.xlsx"}
    )

@router.get("/campaigns/{campaign_id}/export/pdf", responses={200: {"content": {"application/pdf": {}}}}, summary="Export Transactions to PDF")
async def export_campaign_pdf(
    campaign_id: str,
    db: Session = Depends(get_db),
    current_user: Dict[str, Any] = Depends(get_verified_user)
):
    campaign = campaign_repo.get_campaign(db=db, identifier=str(campaign_id))
    if not campaign:
        raise HTTPException(status_code=404, detail="Campaign not found.")
    _verify_group_ownership(db, str(campaign.group_id), current_user.get('sub'))
    
    settings = campaign.settings_override or {}
    
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet
    except ImportError:
        raise HTTPException(status_code=500, detail="reportlab is not installed. Please install it to export PDF files.")
        
    transactions = db.query(Transaction).filter(Transaction.campaign_id == str(campaign.campaign_id), Transaction.status == "approved").order_by(Transaction.created_at.desc()).all()
    
    stream = io.BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=letter)
    elements = []
    
    styles = getSampleStyleSheet()
    
    # Header with KapuLetu Branding if not removed
    if not settings.get("remove_watermark", False):
        branding_header = Paragraph("<font color='green'><b>KapuLetu Campaigns</b></font>", styles['Normal'])
        elements.append(branding_header)
        elements.append(Spacer(1, 6))
        
    title = Paragraph(f"<b>{campaign.title} - Contributions Report</b>", styles['Title'])
    elements.append(title)
    elements.append(Spacer(1, 12))
    
    data = [["Date", "Name", "Phone", "Amount (KES)", "Payment Method"]]
    for txn in transactions:
        data.append([
            txn.created_at.strftime("%Y-%m-%d %H:%M"),
            txn.sender_name or "Anonymous",
            txn.sender_phone or "",
            f"{float(txn.amount):,.2f}",
            txn.payment_method
        ])
        
    t = Table(data)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1A5D1A")), # Kapuletu Green-ish
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor("#F0FDF4")), # Light green tint
        ('GRID', (0, 0), (-1, -1), 1, colors.HexColor("#DDDDDD"))
    ]))
    
    elements.append(t)
    
    # Footer Branding
    if not settings.get("remove_watermark", False):
        elements.append(Spacer(1, 24))
        branding = Paragraph("<i>Report securely generated by KapuLetu - The ultimate community fund manager.</i>", styles['Italic'])
        elements.append(branding)
        
    doc.build(elements)
    
    stream.seek(0)
    
    return StreamingResponse(
        stream,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=campaign_{campaign.slug}_contributions.pdf"}
    )

@router.post("/public/campaigns/{campaign_id}/verify", response_model=CampaignOut, summary="Verify Public Access PIN")
async def public_verify_campaign(
    campaign_id: UUID,
    pin: str = Query(...),
    db: Session = Depends(get_db)
):
    campaign = campaign_repo.get_campaign(db=db, identifier=str(campaign_id))
    if not campaign:
        raise HTTPException(status_code=404, detail="Campaign not found.")
        
    settings = campaign.settings_override or {}
    if settings.get("require_pin", True):
        access_pin = settings.get("access_pin")
        if access_pin and pin != access_pin:
            raise HTTPException(status_code=403, detail="Invalid PIN.")
            
    # If successful, we return the campaign details for the public view
    return campaign
